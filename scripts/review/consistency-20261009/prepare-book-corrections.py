"""Create paragraph-bound corrections for the two verified October manuscripts."""
from pathlib import Path
import json,hashlib,re,argparse
from datetime import date
from docx import Document
ROOT=Path(__file__).resolve().parents[3]
parser=argparse.ArgumentParser();parser.add_argument('--source-dir',default=str(ROOT/'source-assets/originals/books-20261001'))
parser.add_argument('--edition-date',default='2026-10-12',type=lambda value:date.fromisoformat(value).isoformat())
args=parser.parse_args()
SOURCE=Path(args.source_dir)
ratio='T-SROI ist nach Rechenstandard v1.1 das Verhältnis des diskontierten, kausal zurechenbaren direkten und separat belegten transformativen Nettonutzens zum diskontierten Ressourceneinsatz in Euro. Nicht monetarisierbare Folgen bleiben im Wirkungsprofil; Datenqualität und Schutzprüfung sind keine Nutzenmultiplikatoren.'
identity='Eine WÖk-ID ist eine stabile Kennung für einen Prüfgegenstand im Masterregister, kein Messwert und nicht automatisch ein validierter Wirkindikator. Die Beziehung lautet MasterItem -> StateVariable -> Indicator -> Observation -> Analysis / Reality Check. Einheit, Population, Raum, Zeitraum, Systemgrenze, Quellenrolle und Prüfstatus werden in der verknüpften Datenarchitektur geführt.'
profile='Der WÖk-Netto-Wirkungsindex verdichtet ein dimensionsgleiches, vorab definiertes Wirkungsprofil nur bei erfüllter Schutzprüfung. Der externe Nationale Wohlfahrtsindex ist eine andere Kennzahl. Fehlende Evidenz ist offen und wird nicht als neutraler Punktwert 0 ausgegeben. Datenqualität ist Prüfbedingung, kein Punktefaktor.'
changes={
 'I':{
  133:profile+' '+ratio,
  323:identity+' Eine Kennung beweist weder Zustandsveränderung noch Zurechnung. [I-L-12]',
  404:'WÖk-Netto-Wirkungsindex (früher WÖk-NWI). '+profile+' [I-AV-4][I-AV-5]',
  426:'WÖk-ID. '+identity+' [I-AV-4]',
  4062:identity+' Eine transparente Verknüpfung ermöglicht Prüfung und Versionierung, ohne die Kennung zur Bewertung zu erklären. [I-K31-1][I-K31-3]',
  4066:'Eine WÖk-ID adressiert den fachlichen Prüfgegenstand. Für Wasser werden etwa Entnahme, Verbrauch, Stress, Wiederverwendung, Abwasserqualität und Gesundheitsfolgen getrennt operationalisiert. Die zugehörigen Zustandsgrößen, Indikatoren und Beobachtungen tragen ihre Einheiten, Systemgrenzen, Quellen, Zeiträume und Qualitätsstände. Die ID selbst misst keinen dieser Zustände. [I-K31-1][I-K31-3]',
  4081:'Die Registerarchitektur verbindet einen Prüfgegenstand mit seinem SDG-/SDG+-Referenzraum. Welche Zustandsgröße beobachtet wird, welcher Indikator dazu passt und welche Beobachtung vorliegt, sind getrennte Festlegungen. Zielbezug und Kennung erzeugen weder Messwert noch Kausalitätsnachweis.',
  4103:'SDG und SDG+ ordnen Referenzräume. NACE klassifiziert wirtschaftliche Tätigkeiten. ESRS und GRI liefern Berichtsanschlüsse. Die WÖk-ID adressiert einen Prüfgegenstand, dessen Zustandsgrößen, Indikatoren und Beobachtungen gesondert verknüpft werden. Keine dieser Zuordnungen erzeugt allein eine Wirkung oder einen Score. [I-K31-1][I-K31-3]',
  4132:identity+' Vergleichbarkeit entsteht erst aus kompatiblen Definitionen und Systemgrenzen. [I-K31-1][I-K31-3]',
  4196:profile+' [I-K32-6]',
  4200:'Ein positiver Profilwert darf nur bei offenem Schutz-Gate als WÖk-Netto-Wirkungsindex ausgewiesen werden. Ein niedriger oder negativer Profilwert ist ein kritischer Befund. Fehlende Datenqualität oder ungeklärte Mindestbedingungen sind dagegen eine offene Prüfung; sie werden nicht als mittlerer oder neutraler NWI behandelt. [I-K32-2][I-K32-6]',
  4205:profile+' '+ratio+' [I-K32-6][I-K32-7]',
  4206:'Die Profilverdichtung und die monetäre Nutzenrechnung bleiben getrennt. Ein gutes Profil belegt keine Transformation; ein hoher Quotient belegt weder Schutz noch vollständige gesellschaftliche Tragfähigkeit. Transformationspfade benötigen eigene Evidenz. [I-K32-6][I-K32-7]',
  4208:'Die Scorecard zeigt Messgrößen, Einheiten, Referenzen, Evidenz und Schutzbedingungen. Der WÖk-Netto-Wirkungsindex verdichtet nur das geeignete geprüfte Profil. IOI und T-SROI rechnen davon getrennt direkte beziehungsweise zusätzlich belegte transformative Nettonutzenströme in Euro.',
  4420:ratio+' Er übernimmt keinen NWI-Punktwert als Geldbetrag und benötigt eine eigenständige Schutzprüfung.',
  4421:'Dieses Kapitel trennt Profil, direkte monetäre Nutzenrechnung und separat belegte transformative Nutzenströme. Es beschreibt keine fertige Steuerfolge oder Fondsregulierung. Maßgeblich für die Rechnung ist der bestehende Rechenstandard v1.1.',
  4422:'KIIs können relevante Zustandsveränderungen beobachten; die Beobachtung allein beweist keine Attribution. Scorecards ordnen Befunde an dokumentierten Maßstäben und Schutzgrenzen ein. Der WÖk-Netto-Wirkungsindex verdichtet das geeignete Profil. Der monetäre IOI bezieht direkten Nettonutzen auf Ressourcen; T-SROI ergänzt separat belegte transformative Nutzenströme in derselben Preisbasis.',
  4423:'Die Fragen bleiben getrennt: Welche Veränderungen und Schutzbedingungen zeigen sich im Profil? Welcher direkte diskontierte Nettonutzen entsteht je Ressourceneuro? Welcher zusätzliche transformative Nutzenstrom ist gesondert belegt? Keine Kennzahl ersetzt die Wirkpfad-, Verteilungs-, Evidenz- oder Schutzprüfung.',
  4430:'SROI kann neben positiven auch negative, längerfristige und systemische Folgen untersuchen. Seine Aussage hängt von der konkreten Systemgrenze, Gegenfaktik, Attribution und Bewertung ab. Die WÖk beansprucht diese Prüfungen nicht als Erfindung. Ihr T-SROI bleibt ein eigener, transparent begrenzter Rechenstandard und kein Nachweis einer allgemeinen methodischen Überlegenheit.',
  4432:ratio+' Ein negatives kritisches Feld oder eine rote Linie blockiert die positive Ausweisung unabhängig vom Geldquotienten.',
  4433:'Der monetäre IOI ist direkter, kausal begrenzter und diskontierter Nettonutzen in Euro je diskontiertem Ressourceneuro. Ein NWI-Punktwert wird nicht zum Euro-Zähler. Die nichtmonetäre Schutzprüfung bleibt erforderlich.',
  4434:'ROI bezeichnet eine finanzielle Rückflussrechnung mit ausdrücklich dokumentierter Definition. SROI bewertet gesellschaftliche Outcomes unter seinen Methodenbedingungen. Der WÖk-Netto-Wirkungsindex verdichtet ein nichtmonetäres Profil. IOI rechnet direkten, T-SROI zusätzlich separat belegten transformativen Nettonutzen je Ressourceneuro.',
  4439:ratio,
  4441:'WÖk-Netto-Wirkungsindex = Summe_i w_i * s_i, nur bei offenem Schutz-Gate. Die Feldscores müssen auf einer kompatiblen dokumentierten Skala liegen; Gewichte und kritische Felder werden vor der Bewertung festgelegt.',
  4442:'IOI_EUR = PV(direkter kausal begrenzter Nutzen - Schäden) / PV(Investitionen + inkrementelle Ressourcen), nur bei erfüllter Schutzprüfung. Der gesamte direkte Nettonutzen steht im Zähler.',
  4443:'T-SROI = PV_N / PV_R, nur bei G = 1. Transformation geht als separat belegter Nutzenstrom ein, nicht als Transformations-, Resilienz- oder Datenqualitätsmultiplikator.',
  4444:'Das Schutz-Gate ist eine Ausweisungsbedingung, kein Rechenfaktor. Eine rote Linie, ein negatives kritisches Feld, unzureichende Evidenz, eine ungeklärte Systemgrenze oder Zurechnung, eine nicht positive Ressourcenbasis oder eine nicht positive konservative Szenariountergrenze blockieren die positive Kennzahl. Blockiert ist kein Ergebnis 0.',
  4445:'Formelkasten 34-2: monetäre Rechnung nach Rechenstandard v1.1',
  4446:'PV_N = Summe_t=1..T [((B_direkt,t + B_transformativ,t) * a_t * (1-d_t) * (1-v_t) - S_t) / (1+r)^t]. PV_R = Summe_t=0..T [(I_t + K_t) / (1+r_K)^t]. T-SROI = PV_N / PV_R, nur bei G = 1.',
  4447:'B_direkt und B_transformativ sind getrennt belegte Nutzenströme, S sind Schäden innerhalb der Bilanzgrenze. I und K sind Investitionen und inkrementelle Ressourcen. Alle Geldwerte haben dieselbe Preisbasis. Attribution a, Deadweight d und Verdrängung v begrenzen den beanspruchten Nutzen; Schäden werden nicht pauschal mit diesen Faktoren gekürzt. Die Anfangsinvestition bei t = 0 wird nicht abgezinst. Die Diskontsätze r und r_K werden begründet. [I-K34-1][I-K34-5]',
  4448:'Die konservative Szenariountergrenze lautet PV_N^L = Summe_t=1..T [((B_direkt,t + B_transformativ,t) * a_t * (1-d_t) * (1-v_t) * (1-u_t) - S_t) / (1+r)^t]. Der Abschlag u kürzt nur Nutzen, nicht Schäden. Die Untergrenze ist keine statistische Konfidenzgrenze. Dieselben Ressourcen werden nicht zugleich im Nenner und als Schaden im Zähler verbucht.',
  4454:'Eine Transformationsprüfung untersucht Diffusion, Standardsetzung, Infrastruktur, veränderte Risiken und Rückkopplungen. Diese Pfade bleiben auch ohne Geldwert relevant. Im T-SROI werden sie nur dann berücksichtigt, wenn ein eigener kausal begrenzter Nutzenstrom in Euro mit Zeitraum, Empfängerkreis, Preisbasis und Unsicherheit belegt ist. Direkter Nutzen und transformative Folge dürfen dieselbe Veränderung nicht doppelt zählen. [I-K34-1][I-K34-5]',
  4471:'T-SROI ist eine begrenzte monetäre Rechnung innerhalb einer breiteren Transformationsprüfung. Ein hoher Quotient allein beweist weder Transformation noch langfristige Zukunftsfähigkeit.',
  4693:'Die Scorecard und ihr FinalScore liefern im WStG-3.0-Entwurf Evidenz und Schutzlogik. Eine Steuerklasse oder konkrete Steuerfolge entsteht erst aus einem demokratisch bestimmten, steuerartspezifischen Tatbestand. Die historische Score-zu-Tarif-Staffel ist ein Simulationsstand. Belastung und aktive Fondsentlastung sind getrennte Geldflüsse; null zusätzliche Belastung senkt den bisherigen Käuferpreis noch nicht. [I-K38-1][I-K38-2]',
  5914:ratio,
 },
 'II':{
  228:ratio+' Die Transformation selbst wird zusätzlich mit Wirkpfad, Verteilung, Resilienz, Rechten und nichtmonetären Befunden geprüft. [I-L-10]',
  305:'WÖk-Netto-Wirkungsindex (früher WÖk-NWI). '+profile+' [I-AV-4][I-AV-5]',
  316:'T-SROI - Transformational Social Return on Investment. '+ratio+' [I-AV-5]',
  2322:identity+' Die Kennung ist weder Score noch Wirkung. [I-K81-5]',
  6007:ratio,
 }
}
hashes={'I':'dea6cc852ba41a56b34fc0a7c7ea5367f3f7c56a607df43a09007f6ee035e643','II':'c3b4164e1f35f2ff00121e3f5be441f6db89ebde7e880ffcdde1350d6a53673f'}
changes['I'].update({
 314:profile+' [I-L-9]',
 325:ratio+' Die breitere Transformationsprüfung betrachtet zusätzlich Pfade, Verteilung, Resilienz, Rechte und nichtmonetäre Folgen. [I-L-10]',
 413:'SROI - Social Return on Investment. Eine etablierte Methode zur Bewertung gesellschaftlicher Outcomes einschließlich negativer und systemischer Folgen innerhalb der jeweiligen Bilanzgrenze. Der WÖk-T-SROI ist ein gesonderter Modellstandard; seine Reichweite beweist keine höhere empirische Messgenauigkeit. [I-AV-5]',
 415:'T-SROI - Transformational Social Return on Investment. '+ratio+' [I-AV-5]',
 853:profile+' '+ratio,
 1006:'Die Wirkungsökonomie betrachtet Investitionen neben ihrer finanziellen Rendite in einer breiteren Wirkungsprüfung. Das Profil hält nichtmonetäre Folgen und Schutzbedingungen fest. Der T-SROI quantifiziert darin nur belegte direkte und transformative Nettonutzenströme in Euro je diskontiertem Ressourceneuro; er beweist nicht allein einen veränderten Systempfad. [I-K3-11]',
 2617:'Der monetäre IOI setzt direkten, kausal begrenzten und diskontierten Nettonutzen in Euro zum diskontierten Ressourceneinsatz in Beziehung. Der T-SROI ergänzt separat belegte transformative Nutzenströme. Ein NWI-Punktwert ist kein Eurobetrag. Beide Quotienten brauchen eine eigene Schutzprüfung und ersetzen weder das Wirkungsprofil noch die Transformationsprüfung. [I-K16-23]',
 2618:'Netto-Wirkung und Transformation bezeichnen verschiedene Befunde. Transformation betrifft geänderte Strukturen und künftige Entscheidungsbedingungen. Der T-SROI ist dagegen ein begrenzter monetärer Quotient, in dem direkter und zusätzlich belegter transformativer Nettonutzen gemeinsam gerechnet werden. Transformation wird daneben mit Wirkpfaden, Verteilung, Resilienz und Schutzrechten untersucht. [I-K16-4][I-K16-5]',
 2625:ratio+' Eine Transformationsprüfung kann relevante nichtmonetäre Befunde ergeben, die im Quotienten nicht erscheinen. [I-K16-5]',
 2636:'Rohdaten, Zustandsgrößen, Indikatoren, Beobachtungen, WÖk-IDs und Analysen bleiben unterscheidbar. Die ID adressiert den Prüfgegenstand; die Datenarchitektur verknüpft seine Messung. Eine Scorecard ordnet Befunde an dokumentierten Maßstäben. Der WÖk-Netto-Wirkungsindex verdichtet das geeignete geprüfte Profil, IOI und T-SROI rechnen getrennt Nettonutzenströme in Euro. Kein einzelner Wert beweist sämtliche Folgen. [I-K16-4][I-K16-5][I-K16-16]',
 2682:profile+' '+ratio+' Der IOI verwendet dieselbe Ressourcenbasis und Schutzlogik, berücksichtigt aber nur den direkten Nettonutzen. [I-K16-4][I-K16-5][I-K16-23]',
 3361:'Kapital kann Resilienz finanzieren oder Schäden und Systemrisiken verstärken. Eine breitere Kapitalprüfung untersucht diese Wirkpfade. Der T-SROI bildet davon nur die belegten monetären direkten und transformativen Nettonutzenströme je Ressourceneuro ab. Er vermischt sie nicht mit finanzieller Rendite oder einem Resilienz- und Datenqualitätsmultiplikator. [I-23-13]',
 4085:'Der Wirkpfad beschreibt Aktivität, Mechanismus, Zustandsveränderung und weitere Reaktionen für bestimmte Gruppen und eine begründete Systemgrenze. Die WÖk-ID adressiert den Prüfgegenstand; passende Zustandsgrößen, Indikatoren und Beobachtungen werden getrennt verknüpft. SDGs, SDG+ und einschlägige Rechte bieten normative Anschlüsse. Mehrere Zielbezüge vervielfachen nicht dieselbe Veränderung. Verschiedene materielle Folgen bleiben eigenständig prüfbar. [E-K31-10][E-K31-11][E-K31-12][E-K31-13][E-K31-14]',
 4123:'Die mit einer WÖk-ID verknüpfte Messarchitektur dokumentiert Datenquelle, Einheit, Systemgrenze, Datenqualitätsklasse, Prüfanforderung und Version. Die Kennung bleibt stabil, während Daten und Operationalisierung geprüft und versioniert werden. Prüfende müssen Annahmen und Grenzen nachvollziehen können. [I-K31-1][I-K31-3]',
 4418:ratio+' Die Transformationsprüfung reicht darüber hinaus: Sie untersucht auch nichtmonetäre Folgen und geänderte Strukturen.',
 4460:'Fünf mögliche Transformationspfade sind Diffusion, Standardsetzung, Infrastruktur, Kapitalumlenkung und Resilienz. Ihre Analyse beschreibt Mechanismen, betroffene Gruppen, Zeithorizonte und Unsicherheit. Für den T-SROI begründet ein solcher Pfad noch keinen Eurobetrag; dafür ist ein eigener zurechenbarer Nutzenstrom erforderlich.',
 4462:'Tabelle 34-2: Transformationspfade und ihre Nachweisbedingungen',
 4464:'Eine Transformationsprüfung untersucht, ob sich die Bedingungen späterer Entscheidungen verändern. Etablierte Evaluation kann solche Folgen ebenfalls untersuchen. Der T-SROI quantifiziert nur die innerhalb seiner Bilanzgrenze zusätzlich belegten monetären Ströme; der Quotient allein beantwortet die Systemfrage nicht.',
 4468:'Nichtkompensation und Reverse Merit Order schützen vor der Verrechnung schwerer Schäden. Das Schutz-Gate wird für IOI und T-SROI eigenständig geprüft. Es ist kein NWI-Multiplikator und kein Faktor im Nutzenzähler. Rote Linien und negative kritische Felder blockieren die positive Ausweisung unabhängig vom Quotienten. [I-K34-4]',
 4469:'Resilienz gehört zur breiteren Transformationsprüfung: Werden Verwundbarkeit und Abhängigkeit verringert, Lernfähigkeit und institutionelle Stabilität gestärkt? Diese Fragen brauchen eigene Befunde. Der T-SROI darf daraus keinen freien Resilienzfaktor erzeugen. Nur ein separat belegter Nutzenstrom in Euro kann in seine Rechnung eingehen.',
 4474:'Für Unternehmen verbindet Impact Controlling Wirkpfad, Beobachtung, Zurechnung, Bewertung und Entscheidung. Das Profil hält materielle monetäre und nichtmonetäre Folgen sowie Schutzbedingungen fest. Der IOI rechnet den direkten, der T-SROI zusätzlich belegten transformativen Nettonutzen in Euro je Ressourceneuro. Unternehmens- oder Markttransformation bleibt eine breitere, eigenständig nachzuweisende Frage.',
 4475:'Für Kapital bleiben Profil, monetäre Nutzenrechnung und Transformation unterscheidbar. Der WÖk-Netto-Wirkungsindex verdichtet das geeignete geprüfte Profil. Der IOI rechnet direkten Nettonutzen je Ressourceneuro; T-SROI ergänzt separat belegte transformative Ströme. Finanzielle Rendite, Verteilung, Risiko und Systemveränderung brauchen daneben eigene Befunde. [I-K34-2][I-K34-6][I-K34-7]',
 4476:'Öffentliche Maßnahmen brauchen einen objektspezifischen Optionsvergleich unter den bestehenden gesetzlichen Prüfvorgaben. IOI und T-SROI können begrenzte monetäre Rechnungen dazu liefern. Eine Förderung oder Steuerfolge folgt weder automatisch aus NWI noch aus FinalScore oder T-SROI. Der WStG-3.0-Entwurf trennt Evidenzprofil und gesetzlich bestimmten steuerartspezifischen Tatbestand.',
 4477:'Ein Wirkungsfonds braucht ein geprüftes Wirkungsprofil, Nichtkompensation und eine dokumentierte Finanzierungs- und Portfoliologik. IOI und T-SROI können den Vergleich belegter Nutzenströme unterstützen. Diffusion, Standardsetzung, Infrastruktur, Kapitalumlenkung und Resilienz werden daneben als eigene Wirkpfade geprüft. Derselbe Budgeteuro und dieselbe Zustandsveränderung werden nicht mehrfach verbucht.',
 4478:'T-SROI ersetzt keine politische Entscheidung, Rechtsprüfung oder demokratische Aushandlung. Er macht belegte monetäre Nutzenströme innerhalb dokumentierter Grenzen vergleichbar. Ein positiver Quotient bestätigt weder sämtliche Transformationsannahmen noch nichtmonetäre Schutzrechte. Das gesamte Wirkungsprofil und die gesonderte Schutzprüfung bleiben sichtbar.',
 4481:'ROI, SROI, WÖk-Netto-Wirkungsindex, IOI und T-SROI haben getrennte Funktionen. Finanzielle Rückflüsse, gesellschaftliche Outcomes, nichtmonetäres Profil und monetäre Nettonutzenrechnung dürfen nicht ineinander umgerechnet werden. Der aktuelle Rechenstandard v1.1 verwendet keine freien Transformations-, Resilienz- oder Datenqualitätsmultiplikatoren.',
 4482:'Kapitel 30 klärt Messung und Erkenntnisgrenzen. Kapitel 31 verbindet stabile Kennungen mit einer getrennten Messarchitektur. Kapitel 32 behandelt Maßstäbe, Scorecards und Profilverdichtung. Kapitel 33 schützt durch Nichtkompensation und Reverse Merit Order. Kapitel 34 trennt IOI und T-SROI als monetäre Rechnungen von der breiteren Transformationsprüfung.',
 4483:'Das Wirkungsprofil zeigt materielle Folgen und Schutzbedingungen. IOI rechnet direkten Nettonutzen, T-SROI zusätzlich separat belegten transformativen Nettonutzen je Ressourceneuro. Keine dieser Kennzahlen beweist allein eine systemische Veränderung.',
 4504:'Kapitel 34 hat die monetäre T-SROI-Rechnung v1.1 von der breiteren Transformationsprüfung abgegrenzt. Dieses Kapitel erklärt nun, wie Wirkungsdaten verfügbar, prüfbar, verknüpfbar und entscheidungsrelevant werden. Datenverfügbarkeit ist dabei kein Nachweis, dass alle Wirkungen bereits gemessen oder kausal geklärt sind.',
 4509:'Der digitale Produktpass speichert produktbezogene Daten. Die WÖk-ID adressiert den Prüfgegenstand und verbindet dessen dokumentierte Messarchitektur. Benchmarks und Scorecards ordnen Befunde; IOI und T-SROI rechnen belegte monetäre Nutzenströme. Wirkpfad, nichtmonetäre Folgen, Schutzprüfung und Rückkopplung machen daraus eine umfassendere Wirkungsinfrastruktur. [I-K35-3]',
 4531:'Produktpässe machen Daten am Produkt verfügbar, Datenräume verbinden Quellen. WÖk-IDs adressieren Prüfgegenstände. Benchmarks und Scorecards bewerten geeignete Befunde. T-SROI rechnet innerhalb seiner Bilanzgrenze belegte direkte und transformative Nutzenströme. Erst definierte Entscheidungen und spätere Überprüfung schließen die Rückkopplung. [I-K35-1][I-K35-2]',
 4539:'Teil V verbindet Messung und Erkenntnisgrenzen, stabile Kennungen und getrennte Messarchitektur, Maßstäbe und Scorecards, Nichtkompensation und Reverse Merit Order, monetäre IOI-/T-SROI-Rechnungen sowie Produktpässe und Datenräume. Damit werden Befunde prüfbar und nutzbar. Verfügbarkeit, Bewertung, rechtliche Folge und tatsächlich eingetretene Wirkung sind weiterhin getrennte Prüfstände.',
 5069:'KIIs können relevante Zustandsveränderungen beobachten; Kausalität wird eigenständig geprüft. Scorecards ordnen Befunde ein. Der WÖk-Netto-Wirkungsindex verdichtet das geeignete geprüfte Profil. IOI rechnet direkten, T-SROI zusätzlich separat belegten transformativen Nettonutzen in Euro je Ressourceneuro. Diese Funktionen bilden keine automatische Umrechnungskette.',
 5070:profile+' '+ratio+' Finanzielle Tragfähigkeit, nichtmonetäre Folgen und geänderte Systempfade bleiben gesonderte Fragen. [I-K44-6][I-K44-9]',
 5072:'ROI bleibt eine finanzielle Rechnung. KIIs beobachten Zustandsgrößen, die Scorecard strukturiert das Profil. IOI und T-SROI können belegte monetäre Nutzenströme vergleichen. Transformationskraft wird daneben mit Mechanismus, Gegenfaktum, Verteilung, Resilienz und Schutzrechten geprüft. Keine Kennzahl ersetzt die anderen Befunde. [I-K44-6][I-K44-7][I-K44-9]',
 5099:'Wirkungscontrolling verbindet Finanzsteuerung mit Fragestellung, Wirkpfad, Daten, Zurechnung, Bewertung und Entscheidung. Das Wirkungsprofil umfasst nichtmonetäre Folgen und Schutzbedingungen. IOI und T-SROI sind darin begrenzte monetäre Rechnungen. Rückkopplung sorgt dafür, dass die Befunde in Entscheidungen und spätere Korrektur eingehen.',
 5350:profile+' '+ratio+' Ob daraus ein veränderter Systempfad entsteht, braucht zusätzlich eine breitere Transformationsprüfung. [Kap. 34; I-K47-4]',
 5910:'Transformation bedeutet eine Veränderung von Strukturen oder künftigen Entscheidungsbedingungen. Sie wird mit Wirkpfad, Evidenz, Verteilung, Resilienz und Schutzrechten untersucht. Netto-Wirkung, WÖk-Netto-Wirkungsindex und der monetäre T-SROI sind davon zu unterscheiden.',
 5862:profile,
 6023:identity,
 4073:'Eine WÖk-ID adressiert einen fachlichen Prüfgegenstand. Zielbezug, Zustandsgröße, Indikator und Beobachtung werden gesondert festgelegt. Ihre Verknüpfung macht einen Befund prüfbar; die Kennung selbst ist weder Messwert noch Norm noch Kausalitätsbeweis.',
 4090:'Eine WÖk-ID ist eine stabile Kennung. In der verknüpften Datenarchitektur werden Wirkungsraum, Familie, Zustandsgrößen, Indikatoren, Einheiten, Datenquellen, Systemgrenzen, Prüfstatus und Version geführt. Diese Merkmale werden nicht sämtlich in die Kennung hineindefiniert. [I-K31-1][I-K31-3]',
})
changes['II'].update({
 217:profile+' [I-L-9]',
 314:changes['I'][413],
 327:identity+' [I-AV-4]',
 702:'Der Wirkungsfonds braucht ein geprüftes Wirkungsprofil und eine gesonderte Finanzierungslogik. IOI vergleicht belegten direkten Nettonutzen je Ressourceneuro; T-SROI ergänzt separat belegte transformative Nutzenströme. Resilienz, Standards, Pfade und nichtmonetäre Folgen werden zusätzlich untersucht. Automatisierungsgewinne und Alterssicherung sind davon getrennte Verteilungs- und Finanzierungsfragen. [I-K59-4][I-K59-5][I-K59-6]',
 706:'Die Tabelle beschreibt eine Mindestlogik und keine fertige Fondsverfassung. Nichtkompensation und Reverse Merit Order verhindern, dass schwere Schäden durch Vorteile verdeckt werden. IOI und T-SROI können begrenzte monetäre Rechnungen liefern; ein hoher Quotient beweist keine Resilienz oder Systemwirkung. [I-K59-5]',
 724:'IOI setzt belegten direkten Nettonutzen in Euro zu diskontierten Ressourcen in Beziehung. T-SROI ergänzt separat belegte transformative Ströme. Ob Systeme resilienter, gerechter, regenerativer oder demokratisch stabiler werden, braucht eigenständige nichtmonetäre Befunde und Schutzprüfung. [I-K59-5]',
 740:'Kapital kann Regeneration, Bildung, Pflege und Infrastruktur ermöglichen oder Schäden und Abhängigkeit verstärken. Wirkungsfonds prüfen deshalb das gesamte Portfolio, sein Wirkungsprofil, rote Linien und die tatsächliche Finanzierung. IOI und T-SROI bilden darin begrenzte monetäre Nutzenrechnungen. Transformations- und Wirkungsrisiken sowie Versicherbarkeit liefern zusätzliche, eigenständig zu prüfende Befunde.',
 6003:changes['I'][5910],
 5955:profile,
 6116:identity,
})
# Historical works remain named and dated. Current calculations cite the actual
# v1.1 standard instead of attributing its formula to a 2025 multiplier paper.
historical_note=' Historische Konzeptquelle; ihre Multiplikatorlogik gilt nicht als aktuelle Rechenregel. Für die aktuelle Rechnung: T-SROI-Rechenstandard v1.1, 2.8.2026, und führender Begriffsleitfaden v1.8, 1.10.2026, §§ 5 und 10.10.'
table_updates={
 'I':{(7,4,1):'Stabile Kennungen für Prüfgegenstände; Zustandsgrößen, Indikatoren, Beobachtungen und Analysen werden gesondert verknüpft.',
 (9,3,1):'Wie ist das vorab definierte Wirkungsprofil unter seinen Schutzbedingungen einzuordnen?',
 (9,3,2):'WÖk-Netto-Wirkungsindex: zulässige Profilverdichtung auf dokumentierter Skala. Evidenz und harte Schutzgrenzen werden separat geprüft.',
 (9,4,1):'Welcher belegte direkte und transformative Nettonutzen entsteht je diskontiertem Ressourceneuro?',
 (9,4,2):'T-SROI = PV_N / PV_R nach v1.1, nur bei offenem Schutz-Gate. Kein NWI-Punktwert, kein freier Multiplikator.',
 (10,1,2):'Übernahme und ihre Folgen brauchen eigene Beobachtungen und Zurechnung. Nur ein zusätzlich belegter monetärer Nutzenstrom geht in T-SROI ein.',
 (10,4,2):'Veränderte Kapitalströme sind ein möglicher Transformationspfad. Ihr monetärer Nutzen wird eigenständig nachgewiesen und nicht mit dem finanzierten Projekt doppelt gezählt.'},
 'II':{}
}
phrase_updates={
 'WÖk-IDs machen Indikatoren adressierbar.':'WÖk-IDs adressieren Prüfgegenstände und verknüpfen deren dokumentierte Messarchitektur.',
 'WÖk-IDs adressieren Indikatoren.':'WÖk-IDs adressieren Prüfgegenstände; die Datenarchitektur verbindet Zustandsgrößen, Indikatoren und Beobachtungen.',
 'WÖk-IDs ordnen Indikatoren.':'WÖk-IDs adressieren Prüfgegenstände und verknüpfen deren dokumentierte Messarchitektur.',
 'WÖk-IDs ordnen Wirkungsindikatoren.':'WÖk-IDs adressieren Prüfgegenstände und verknüpfen deren dokumentierte Messarchitektur.',
 'WÖk-IDs machen sie methodisch adressierbar.':'WÖk-IDs adressieren Prüfgegenstände und verknüpfen deren dokumentierte Messarchitektur.',
 'WÖk-IDs machen die einzelnen Indikatoren adressierbar.':'WÖk-IDs adressieren Prüfgegenstände und verknüpfen die passenden Indikatoren gesondert.',
}
changes['I'].update({
 69:'Kapitel 34 - T-SROI und die Prüfung von Transformationspfaden\t539',
 4416:'Kapitel 34 - T-SROI und die Prüfung von Transformationspfaden',
 2239:'Das Wirkungssteuergesetz ist dafür ein Beispiel: Ein gesetzlich bestimmtes Modul kann auf geeignete Daten und Nachweise zurückgreifen und konkrete wirtschaftliche Folgen auslösen. Scorecard, Prüfstatus und gesetzlicher Tatbestand bleiben getrennt [I-12-8][I-12-9]. T-SROI liefert dazu eine begrenzte monetäre Nutzenrechnung nach Rechenstandard v1.1. Resilienz, nicht monetarisierbare Folgen, Verteilung und Datenqualität werden eigenständig geprüft; sie sind keine frei gewählten Faktoren des Quotienten.',
 2707:'[I-K16-23] Eigene methodische Definition nach Rechenstandard v1.1: IOI_EUR setzt direkten, kausal begrenzten und diskontierten Nettonutzen in Euro zu diskontierten Ressourcen in Euro in Beziehung. Profilverdichtung, finanzielle Rendite und zusätzlich belegte transformative Nutzenströme bleiben davon getrennt.',
 3003:'WÖk-IDs, Scorecards, Benchmarks und digitale Produktpässe verbinden Prüfgegenstände mit ihren Daten und Bewertungsmaßstäben [I-19-15]. T-SROI ergänzt eine begrenzte monetäre Rechnung: direkte und separat belegte transformative Nutzenströme, kausale Begrenzungen, Schäden und Ressourcen werden in derselben Preisbasis verglichen. Resilienz und nicht monetarisierbare Folgen bleiben gesonderte Befunde; Datenqualität ist Prüfbedingung.',
 3075:'Kapitalrückkopplung braucht mehr als finanzielle Rentabilität. Wirkpfad, Verteilung, Zeitverzug, Risiken, Resilienz und Schutzrechte werden eigenständig geprüft [I-K20-5]. T-SROI kann innerhalb einer dokumentierten Bilanzgrenze belegten direkten und transformativen Nettonutzen je Ressourceneuro vergleichen. Erst die Verbindung der getrennten Befunde mit einer Entscheidung und späterer Überprüfung bildet die Rückkopplung.',
 4236:'Transformationswirkung ist kontextabhängig. Dieselbe Investition kann in unterschiedlichen Systemzuständen andere Folgen entfalten. Eine Solarinvestition in einem fossil geprägten Energiesystem hat andere mögliche Wirkpfade als dieselbe Investition in einem bereits weitgehend regenerativen System. Bei Bildung zählen auch institutionelle Bedingungen und der demokratische Resonanzraum. Diese Unterschiede benötigen eigene Evidenz und Zurechnung; sie werden nicht durch einen frei gewählten Transformationsfaktor in T-SROI ersetzt.',
 4398:'Das nächste Kapitel trennt die monetäre T-SROI-Rechnung von der breiteren Prüfung von Transformationspfaden. Nichtkompensation und Reverse Merit Order bleiben Schutzbedingungen: Ein hoher Geldquotient darf Schäden, Rechte und kritische Verteilungsfolgen nicht verdecken.',
 5074:'Wirkungscontrolling führt getrennte Befunde in die Entscheidung ein: Wirkpfad, beobachtete oder erwartete Zustandsveränderungen, Wirkungsprofil, Schutzbedingungen, Verteilung, Evidenz und Unsicherheit. IOI_EUR und T-SROI liefern dazu begrenzte monetäre Rechnungen nach v1.1. Sie werden neben nicht monetarisierbaren Folgen und der breiteren Transformationsprüfung ausgewiesen. Anschließend wird überprüft, welche Folgen tatsächlich eintreten; eine Kennzahl legitimiert kein Projekt nachträglich.',
})
for v,indices in {'I':[2689,4489,5112],'II':[747]}.items():
    source_doc=Document(SOURCE/f'PRINT_6x9_Band-{v}_SDGplus_2026-10-01_Entwurf.docx')
    for i in indices:changes[v][i]=source_doc.paragraphs[i].text+historical_note
entries=[]
for v,updates in changes.items():
    name=f'PRINT_6x9_Band-{v}_SDGplus_2026-10-01_Entwurf.docx';path=SOURCE/name
    if hashlib.sha256(path.read_bytes()).hexdigest()!=hashes[v]:raise RuntimeError('Manuscript changed: '+name)
    doc=Document(path)
    for i,p in enumerate(doc.paragraphs):
        if i in updates:continue
        if p.text.startswith('[') and ('Whitepaper T-SROI' in p.text or 'T-SROI - Der neue Standard' in p.text):
            updates[i]=p.text+historical_note
        else:
            new=p.text
            for before,after in phrase_updates.items():new=new.replace(before,after)
            if new!=p.text:updates[i]=new
    for i,new in updates.items():
        p=doc.paragraphs[i]
        if p._p.xpath('.//w:drawing'):raise RuntimeError('Drawing must be handled separately')
        entries.append({'source':name,'sourceSha256':hashes[v],'volume':v,'paragraphIndex':i,'beforeText':p.text,'afterText':new,'basis':'Begriffsleitfaden v1.8 §§ 5, 10.10; T-SROI-Rechenstandard v1.1; WStG 3.0 §§ 55-61, 66, 100'})
    for (t,r,c),new in table_updates[v].items():
        entries.append({'source':name,'sourceSha256':hashes[v],'volume':v,'tableIndex':t,'rowIndex':r,'cellIndex':c,'beforeText':doc.tables[t].rows[r].cells[c].text,'afterText':new,'basis':'Begriffsleitfaden v1.8 §§ 5, 10.10; T-SROI-Rechenstandard v1.1'})
def key(e):
    return (e['volume'],e['paragraphIndex']) if 'paragraphIndex' in e else (e['volume'],e['tableIndex'],e['rowIndex'],e['cellIndex'])
combined={key(e):e for e in entries}
additional=ROOT/'content/grundlagen/book-corrections-additional-20261010.json'
if additional.exists():
    for e in json.loads(additional.read_text())['entries']:
        doc=Document(SOURCE/e['source'])
        actual=doc.paragraphs[e['paragraphIndex']].text if 'paragraphIndex' in e else doc.tables[e['tableIndex']].rows[e['rowIndex']].cells[e['cellIndex']].text
        if hashlib.sha256((SOURCE/e['source']).read_bytes()).hexdigest()!=e['sourceSha256'] or actual!=e['beforeText']:
            raise RuntimeError('Unknown additional correction source')
        if key(e) in combined and combined[key(e)]['beforeText']!=e['beforeText']:
            raise RuntimeError('Conflicting correction provenance')
        combined[key(e)]=e
entries=list(combined.values())
save=ROOT/'content/grundlagen/book-corrections-20261009.json';save.write_text(json.dumps({'date':args.edition_date,'status':'source-bound correction set; full publication review pending','entries':entries},ensure_ascii=False,indent=2)+'\n')
print(f'{len(entries)} explicit paragraph corrections prepared against verified original manuscripts.')
