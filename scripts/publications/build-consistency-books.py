"""Rebuild the October manuscripts from checksum-bound, reviewable corrections.

Original manuscripts are immutable. Unknown paragraphs or cells fail closed.
Render the resulting DOCX with the bundled documents renderer before publishing.
"""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
import hashlib,json,argparse,re
from datetime import date
from lxml import etree
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.run import Run
from docx.text.paragraph import Paragraph
import pymupdf
import pyphen
from book_figure_integrity import validate_figure
from book_typography import format_umlaut_ligatures
from book_source_edit import replace_paragraph as replace

ROOT=Path(__file__).resolve().parents[2]
parser=argparse.ArgumentParser()
parser.add_argument('--source-dir',default=str(ROOT/'source-assets/originals/books-20261001'))
parser.add_argument('--output-dir',default=str(ROOT/'source-assets/generated/books-20261009'))
parser.add_argument('--table-font-regular',type=Path)
parser.add_argument('--table-font-bold',type=Path)
parser.add_argument('--volume',choices=['I','II'])
args=parser.parse_args();source=Path(args.source_dir);out=Path(args.output_dir);out.mkdir(parents=True,exist_ok=True)
font_root=Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/native/libreoffice-headless'
def table_font(argument,name):
    matches=[argument] if argument else list(font_root.rglob(name))
    if len(matches)!=1 or not matches[0].is_file():
        raise RuntimeError('Supply the actual regular and bold table fonts for reproducible width measurement')
    return matches[0]
regular_font=table_font(args.table_font_regular,'FrankRuhlHofshi-Regular.otf')
bold_font=table_font(args.table_font_bold,'FrankRuhlHofshi-Bold.otf')
corrections=json.loads((ROOT/'content/grundlagen/book-corrections-20261009.json').read_text())
edition_date=date.fromisoformat(corrections['date'])
months=('Januar','Februar','März','April','Mai','Juni','Juli','August','September','Oktober','November','Dezember')
edition_label=f'{edition_date.day}. {months[edition_date.month-1]} {edition_date.year}'
figure_sources=json.loads((ROOT/'content/grundlagen/book-figure-corrections-20261010.json').read_text())['entries']
figure_build={e['key']:e for e in json.loads((ROOT/'reports/consistency-20261009/book-figure-build.json').read_text())}

def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def public_text(text):return re.sub('[\u2010\u2011\u2012\u2013\u2014\u2212]','-',text).replace('\u00ad','')
def images(path):
    with ZipFile(path) as z:return {n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
def mathematical_glyphs(doc):
    """Use ordinary mathematical glyphs instead of the Hebrew font's alternatives.

    Keep every character and the surrounding original typography unchanged.
    The renderer substitutes Frank Ruhl Hofshi for some unavailable bold faces;
    its plus/greater-than glyphs are visually ambiguous in scores and ID chains.
    """
    changed=0
    for node in list(doc.element.xpath('.//w:r')):
        run=Run(node,None);original=run.text
        if not re.search(r'[+<>Σ∑×→≥≤]',original):continue
        replacements=[]
        # Preserve tabs, breaks, discretionary hyphens and embedded objects in
        # their exact order. A rendered-page-break hint must not make a caption
        # ineligible for the same mathematical typography as the table cells.
        for child in list(node):
            if child.tag==qn('w:rPr'):continue
            parts=re.split(r'([+<>Σ∑×→≥≤])',child.text or '') if child.tag==qn('w:t') else [None]
            for part in parts:
                if part=='':continue
                copy=deepcopy(node)
                for item in list(copy):
                    if item.tag!=qn('w:rPr'):copy.remove(item)
                new_run=Run(copy,None)
                if part is None:copy.append(deepcopy(child))
                else:new_run.add_text(part)
                if part is not None and len(part)==1 and re.fullmatch(r'[+<>Σ∑×→≥≤]',part):
                    fonts=new_run._r.get_or_add_rPr().get_or_add_rFonts()
                    for name in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+name),'DejaVu Sans')
                    changed+=1
                node.addprevious(copy);replacements.append(new_run.text)
        if ''.join(replacements)!=original:raise RuntimeError('Mathematical font repair changed text')
        node.getparent().remove(node)
    return changed

def format_reference_breaks(doc,volume):
    """Keep the two visually reviewed references with their final URLs."""
    if volume!='I':return []
    expected={
        1860:('E-K8-8','[E-K8-8] European Union: Corporate Sustainability Reporting Directive,','corporate-sustainability-reporting_en'),
        4414:('E-K33-2','[E-K33-2] Meadows, Donella H.: Thinking in Systems.','systems-thinking-resources/'),
    }
    records=[]
    for index,(reference,prefix,suffix) in expected.items():
        paragraph=doc.paragraphs[index];text=paragraph.text
        if not text.startswith(prefix) or suffix not in text:
            raise RuntimeError('Unexpected source of the reviewed reference-page break')
        paragraph.paragraph_format.keep_together=True
        if paragraph.text!=text:raise RuntimeError('Reference pagination changed source text')
        records.append({'paragraphIndex':index,'reference':reference,'keepTogether':True,
                        'sourceTextUnchanged':True,'sourceTextSha256':hashlib.sha256(text.encode()).hexdigest()})
    return records

def format_sdg_explanation_paragraphs(doc,volume):
    """Use body style for the five corrected explanations, preserving their text.

    The source slots were short Heading 3 items. Their complete corrected
    explanations must not inherit enlarged heading type and keep-with-next
    across several pages. The actual subsection headings remain headings.
    """
    if volume!='I':return []
    expected={
        3860:'Die SDGs und SDG+ stehen deshalb nicht allein.',
        3861:'Diese Ebenen werden nicht zu einer scheinbar einfachen Rechts-Pyramide',
        3862:'Darum beginnt jede belastbare Einordnung mit dem Wirkpfad.',
        3863:'Die methodische Regel lautet:',
        3864:'Ein einfaches Beispiel ist die Vergabe eines kommunalen Mittagessens.',
    }
    records=[]
    for index,prefix in expected.items():
        paragraph=doc.paragraphs[index];text=paragraph.text
        if paragraph.style.name!='Heading 3' or not text.startswith(prefix):
            raise RuntimeError('Unknown source of the reviewed SDG body-style correction')
        paragraph.style=doc.styles['Normal']
        paragraph.paragraph_format.keep_with_next=False
        if paragraph.text!=text:raise RuntimeError('SDG body-style correction changed source text')
        records.append({'paragraphIndex':index,'beforeStyle':'Heading 3','afterStyle':'Normal',
                        'keepWithNext':False,'sourceTextUnchanged':True,
                        'sourceTextSha256':hashlib.sha256(text.encode()).hexdigest()})
    return records

def format_short_paragraph_continuations(doc,volume):
    """Keep the two reviewed short page continuations with their paragraphs."""
    if volume!='I':return []
    expected={
        4117:('Datenqualität beschreibt, wie verlässlich, vollständig, aktuell, granular und prüfbar ein Wert ist.', '[I-K31-1].'),
        4189:('Eine Scorecard enthält damit mindestens fünf Ebenen.', '[I-K32-1][E-K32-1].'),
    }
    records=[]
    for index,(prefix,suffix) in expected.items():
        paragraph=doc.paragraphs[index];text=paragraph.text
        if not text.startswith(prefix) or not text.endswith(suffix):
            raise RuntimeError('Unknown source of the reviewed short paragraph continuation')
        paragraph.paragraph_format.keep_together=True
        if paragraph.text!=text:raise RuntimeError('Paragraph pagination changed source text')
        records.append({'paragraphIndex':index,'keepTogether':True,'sourceTextUnchanged':True,
                        'sourceTextSha256':hashlib.sha256(text.encode()).hexdigest()})
    lead,closing=doc.paragraphs[5551],doc.paragraphs[5552]
    lead_text,closing_text=lead.text,closing.text
    if not lead_text.startswith('Damit bereitet Kapitel 50 das konkrete Beispiel vor.') or closing_text!='Diese Frage führt zu [Kap. 51]: Das Apfelbeispiel.':
        raise RuntimeError('Unknown source of the reviewed isolated chapter closing')
    lead.paragraph_format.keep_with_next=True
    closing.paragraph_format.keep_together=True
    if lead.text!=lead_text or closing.text!=closing_text:
        raise RuntimeError('Closing pagination changed source text')
    records.append({'paragraphIndex':5551,'keepWithNext':True,'nextParagraphIndex':5552,
                    'sourceTextUnchanged':True,'sourceTextSha256':hashlib.sha256(lead_text.encode()).hexdigest()})
    return records

def format_tables(doc):
    """Keep every cell and comparison dimension; separate the printed columns."""
    formatted=[];font=pymupdf.Font(fontfile=str(regular_font));header_font=pymupdf.Font(fontfile=str(bold_font));hyphenation=pyphen.Pyphen(lang='de_DE',left=3,right=3)
    for index,table in enumerate(doc.tables):
        if len(table.columns)<2:continue
        previous=table._tbl.getprevious();caption=None
        if previous is not None and previous.tag==qn('w:p'):
            paragraph=Paragraph(previous,doc)
            if re.match(r'^Tabelle \d+-\d+: ',paragraph.text):
                text=paragraph.text
                paragraph.paragraph_format.keep_with_next=True
                if paragraph.text!=text:raise RuntimeError('Table-caption pagination changed source text')
                caption={'text':text,'keepWithNext':True,'sourceTextUnchanged':True}
        before=[[cell.text for cell in row.cells] for row in table.rows]
        total=sum(column.width.pt for column in table.columns)
        widths=[column.width.pt for column in table.columns]
        minimum=[max((header_font.text_length(word,fontsize=9) for word in cell.text.split()),default=0)+11
                 for cell in table.rows[0].cells]
        if sum(minimum)>total:raise RuntimeError(f'Table {index} needs individual header layout review')
        demand=[max(widths[column],max((font.text_length(word,fontsize=9)+9
                    for row in table.rows for word in row.cells[column].text.split()),default=0))
                for column in range(len(widths))]
        needs=[max(0,w-floor) for w,floor in zip(demand,minimum)]
        extra=total-sum(minimum);weight=sum(needs)
        widths=[floor+extra*(need/weight if weight else 1/len(widths)) for floor,need in zip(minimum,needs)]
        discretionary_breaks=0
        table.autofit=False
        for column,width in zip(table.columns,widths):column.width=Pt(width)
        for row_index,row in enumerate(table.rows):
            props=row._tr.get_or_add_trPr()
            if props.find(qn('w:cantSplit')) is None:props.append(OxmlElement('w:cantSplit'))
            if row_index==0 and props.find(qn('w:tblHeader')) is None:props.append(OxmlElement('w:tblHeader'))
            for cell,width in zip(row.cells,widths):
                cell.width=Pt(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
                props=cell._tc.get_or_add_tcPr();margins=props.find(qn('w:tcMar'))
                if margins is None:margins=OxmlElement('w:tcMar');props.append(margins)
                for edge,value in [('top','50'),('bottom','50'),('left','70'),('right','70')]:
                    node=margins.find(qn('w:'+edge))
                    if node is None:node=OxmlElement('w:'+edge);margins.append(node)
                    node.set(qn('w:w'),value);node.set(qn('w:type'),'dxa')
                for paragraph in cell.paragraphs:
                    paragraph.alignment=WD_ALIGN_PARAGRAPH.LEFT
                    paragraph.paragraph_format.space_before=Pt(0)
                    paragraph.paragraph_format.space_after=Pt(3)
                    for run in paragraph.runs:
                        run.font.size=Pt(9)
                        run.font.name='Frank Ruhl Hofshi'
                        # OOXML discretionary hyphens leave the cell's text
                        # unchanged and print only at a genuine line break.
                        original=run.text
                        if not original or any(child.tag not in {qn('w:rPr'),qn('w:t')} for child in run._r):continue
                        parts=re.split(r'([A-Za-zÄÖÜäöüß]{10,})',original)
                        run_font=header_font if run.bold else font
                        if not any(len(word)>=10 and run_font.text_length(word,fontsize=9)>width-12 for word in parts):continue
                        run.clear()
                        for word in parts:
                            positions=hyphenation.positions(word) if re.fullmatch(r'[A-Za-zÄÖÜäöüß]{10,}',word) and run_font.text_length(word,fontsize=9)>width-12 else []
                            start=0
                            for stop in [*positions,len(word)]:
                                if stop>start:
                                    text=OxmlElement('w:t');text.set(qn('xml:space'),'preserve');text.text=word[start:stop];run._r.append(text)
                                if stop<len(word):run._r.append(OxmlElement('w:softHyphen'));discretionary_breaks+=1
                                start=stop
                        if run.text!=original:raise RuntimeError('Hyphenation changed table text')
        if before!=[[cell.text for cell in row.cells] for row in table.rows]:raise RuntimeError('Table formatting changed content')
        formatted.append({'tableIndex':index,'columnWidthsPt':[round(w,2) for w in widths],'fontPt':9,'fontFamily':'Frank Ruhl Hofshi','regularFontSha256':sha(regular_font),'boldFontSha256':sha(bold_font),'cellMarginPt':3.5,'splitRows':False,'discretionaryBreaks':discretionary_breaks,'hyphenationDictionary':'Pyphen 0.17.2 / de_DE','captionLayout':caption})
    return formatted

def format_empty_page_breaks(doc,volume):
    """Retain explicit breaks without spilling their empty line onto a new page."""
    changed=[]
    for index in [103,151,3814,4111,4171,4415] if volume=='I' else []:
        paragraph=doc.paragraphs[index]
        breaks=paragraph._p.xpath('.//w:br[@w:type="page"]')
        if paragraph.text.strip() or len(breaks)!=1 or paragraph._p.xpath('.//w:drawing'):
            raise RuntimeError(f'Unknown empty page-break source: {volume}/{index}')
        paragraph.paragraph_format.space_before=Pt(0)
        paragraph.paragraph_format.space_after=Pt(0)
        paragraph.paragraph_format.line_spacing=Pt(1)
        paragraph.paragraph_format.keep_with_next=False
        for run in paragraph.runs:run.font.size=Pt(1)
        change={'paragraphIndex':index,'explicitPageBreakPreserved':True,'emptyLinePt':1}
        if index==151:
            prior_image=doc.paragraphs[index-1]
            next_heading=doc.paragraphs[index+1]
            drawings=prior_image._p.xpath('.//w:drawing')
            if (prior_image.text.strip() or len(drawings)!=1
                    or doc.paragraphs[index-2].text!='Das Modell auf einen Blick'
                    or next_heading.text!='Vorspann'):
                raise RuntimeError('Unknown page-break context after the introductory model figure')
            drawing_xml=etree.tostring(drawings[0])
            breaks[0].getparent().remove(breaks[0])
            before=prior_image.text
            run=prior_image.add_run();run.font.size=Pt(1)
            page_break=OxmlElement('w:br');page_break.set(qn('w:type'),'page');run._r.append(page_break)
            paragraph.paragraph_format.keep_with_next=True
            if prior_image.text!=before or etree.tostring(drawings[0])!=drawing_xml:
                raise RuntimeError('Introductory page-break relocation changed text or drawing')
            change['pageBreakRelocatedToParagraphEnd']=index-1
            change['imageDrawingPreserved']=True
        elif index==3814:
            next_heading=doc.paragraphs[index+1]
            if next_heading.text!='Exkurs: Warum die SDGs der Referenzrahmen der Wirkungsökonomie sind' or doc.paragraphs[index-1].text.strip():
                raise RuntimeError('Unknown SDG-excursus page-break context')
            breaks[0].getparent().remove(breaks[0])
            prior_paragraph=doc.paragraphs[index-2]
            if not prior_paragraph.text.startswith('[E-K28-7] Tocqueville, Alexis de:'):
                raise RuntimeError('Unknown paragraph preceding the SDG excursus')
            before=prior_paragraph.text
            run=prior_paragraph.add_run();run.font.size=Pt(1)
            page_break=OxmlElement('w:br');page_break.set(qn('w:type'),'page');run._r.append(page_break)
            if prior_paragraph.text!=before:raise RuntimeError('Page-break relocation changed source text')
            for spacer in [doc.paragraphs[index-1],paragraph]:
                spacer.paragraph_format.space_before=Pt(0)
                spacer.paragraph_format.space_after=Pt(0)
                spacer.paragraph_format.line_spacing=Pt(1)
                spacer.paragraph_format.keep_with_next=True
            change['pageBreakRelocatedToParagraphEnd']=index-2
        elif index==4111:
            prior_paragraph=doc.paragraphs[index-1]
            caption=doc.paragraphs[index+1]
            if (not prior_paragraph.text.startswith('Drittens bereiten Indikatorfamilien spätere Bewertung vor.')
                    or caption.text!='Tabelle 31-2: Beispiele für Indikatorfamilien'):
                raise RuntimeError('Unknown source of the blank page before the indicator-family table')
            breaks[0].getparent().remove(breaks[0])
            before=prior_paragraph.text
            run=prior_paragraph.add_run();run.font.size=Pt(1)
            page_break=OxmlElement('w:br');page_break.set(qn('w:type'),'page');run._r.append(page_break)
            paragraph.paragraph_format.keep_with_next=True
            if prior_paragraph.text!=before:raise RuntimeError('Table-page-break relocation changed source text')
            change['pageBreakRelocatedToParagraphEnd']=index-1
        changed.append(change)
    return changed

report_path=ROOT/'reports/consistency-20261009/book-source-build.json'
report=[entry for entry in json.loads(report_path.read_text()) if entry['volume']!=args.volume] if args.volume and report_path.exists() else []
for volume in [args.volume] if args.volume else ['I','II']:
    entries=[x for x in corrections['entries'] if x['volume']==volume]
    path=source/entries[0]['source']
    if sha(path)!=entries[0]['sourceSha256']:raise RuntimeError('Source checksum mismatch: '+str(path))
    doc=Document(path)
    original_paragraphs=len(doc.paragraphs);original_tables=len(doc.tables);original_sections=len(doc.sections)
    changed_images={}
    for figure in figure_sources:
        built=validate_figure(ROOT,figure,figure_build[figure['key']]);replacement=ROOT/built['image']
        if sha(replacement)!=built['sha256']:raise RuntimeError('Rendered figure checksum mismatch')
        for placement in figure['placements']:
            if placement['volume']!=volume:continue
            media=placement['media'];parts={r.target_part.partname.rsplit('/',1)[-1]:r.target_part for r in doc.part.rels.values() if r.reltype.endswith('/image')}
            if media not in parts:raise RuntimeError('Image relationship missing: '+media)
            part=parts[media]
            if hashlib.sha256(part.blob).hexdigest()!=figure['originalSha256']:raise RuntimeError('Unknown original image: '+media)
            if replacement.suffix!=Path(media).suffix:raise RuntimeError('Image MIME/extension changed')
            part._blob=replacement.read_bytes()
            for shape in doc.inline_shapes:
                blips=shape._inline.xpath('.//a:blip')
                for blip in blips:
                    rid=blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                    if rid and doc.part.rels[rid].target_part==part:
                        shape.height=int(shape.width*built['size'][1]/built['size'][0])
                        shape._inline.docPr.set('descr',figure['title']+'. '+figure['subtitle']+'. '+figure['note'])
            changed_images['word/media/'+media]=built['sha256']
    for e in entries:
        if 'paragraphIndex' in e:p=doc.paragraphs[e['paragraphIndex']]
        else:
            cell=doc.tables[e['tableIndex']].rows[e['rowIndex']].cells[e['cellIndex']]
            if cell.text!=e['beforeText']:raise RuntimeError('Unknown table cell')
            if len(cell.paragraphs)!=1:raise RuntimeError('Multi-paragraph cell requires explicit treatment')
            p=cell.paragraphs[0]
        if p.text!=e['beforeText']:raise RuntimeError('Unknown paragraph: '+str(e.get('paragraphIndex')))
        prefix_image={('I',687):'image5.png',('II',2777):'image66.jpeg'}.get((volume,e.get('paragraphIndex')))
        replace(p,e['afterText'],prefix_image=prefix_image)
    # Update the existing stand line, rather than insert a substitute preface.
    stand_changes=[]
    for i,p in enumerate(doc.paragraphs[:30]):
        if '1. Oktober 2026' in p.text or '01.10.2026' in p.text or '1.10.2026' in p.text:
            before=p.text;new=f'Lesefassung vom {edition_label}. Wirtschafts- und Gesellschaftsmodell (Entwurf). Begriffsgrundlage: Begriffsleitfaden v1.8. Rechenstandard: T-SROI v1.1. Die historischen Manuskripte bleiben erhalten.'
            if volume=='II':new+=' Die gemeinsamen Belege zu Kurzfassung, Begriffen, Vorwort und Abkürzungen stehen in Band I unter denselben Quellenkennungen.'
            replace(p,new);stand_changes.append({'paragraphIndex':i,'beforeText':before,'afterText':new})
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    doc.core_properties.author='Natalie Weber'
    doc.core_properties.subject=f'Wirtschafts- und Gesellschaftsmodell; Lesefassung vom {edition_label}'
    typography_changes=0
    for t in doc.element.xpath('.//w:t'):
        if t.text and public_text(t.text)!=t.text:
            t.text=public_text(t.text);typography_changes+=1
    table_layout=format_tables(doc)
    page_break_layout=format_empty_page_breaks(doc,volume)
    reference_break_layout=format_reference_breaks(doc,volume)
    sdg_body_layout=format_sdg_explanation_paragraphs(doc,volume)
    short_continuation_layout=format_short_paragraph_continuations(doc,volume)
    math_glyph_changes=mathematical_glyphs(doc)
    umlaut_ligature_layout=format_umlaut_ligatures(doc)
    target=out/f'PRINT_6x9_Band-{volume}_Konsistenz_2026-10-10.docx';doc.save(target)
    check=Document(target)
    assert len(check.paragraphs)==original_paragraphs and len(check.tables)==original_tables and len(check.sections)==original_sections
    expected=images(path);expected.update(changed_images)
    assert images(target)==expected,'Unreviewed illustration change'
    for e in entries:
        actual=check.paragraphs[e['paragraphIndex']].text if 'paragraphIndex' in e else check.tables[e['tableIndex']].rows[e['rowIndex']].cells[e['cellIndex']].text
        assert actual==public_text(e['afterText'])
    report.append({'volume':volume,'source':str(path.relative_to(ROOT)) if path.is_relative_to(ROOT) else path.name,'sourceSha256':sha(path),'output':str(target.relative_to(ROOT)) if target.is_relative_to(ROOT) else target.name,'outputSha256':sha(target),'corrections':len(entries),'paragraphs':original_paragraphs,'tables':original_tables,'sections':original_sections,'preservedImages':len(images(path))-len(changed_images),'correctedImages':changed_images,'standChanges':stand_changes,'typographyRunsNormalized':typography_changes,'mathematicalGlyphRuns':math_glyph_changes,'umlautLigatureLayout':umlaut_ligature_layout,'tableLayout':table_layout,'pageBreakLayout':page_break_layout,'referenceBreakLayout':reference_break_layout,'sdgBodyLayout':sdg_body_layout,'shortContinuationLayout':short_continuation_layout})
save=ROOT/'reports/consistency-20261009/book-source-build.json';save.parent.mkdir(parents=True,exist_ok=True);save.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(report,ensure_ascii=False,indent=2))
