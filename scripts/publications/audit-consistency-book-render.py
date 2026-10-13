"""Bind the page-review inventory to the actual rebuilt book PDFs."""
from pathlib import Path
import argparse,json,hashlib,re,bisect
import pymupdf
from docx import Document

ROOT=Path(__file__).resolve().parents[2]
def normalize(text):
    text=re.sub(r'\n\s*\d+\s*$','',text)
    return re.sub(r'[^a-z0-9äöüß]','',text.lower().replace('\u00ad',''))
def signed_normalize(text):
    # Preserve signed numeric values while ignoring typesetting hyphenation.
    text=re.sub(r'([+\-−])\s*(?=\d)',lambda m:'POSNUM' if m[1]=='+' else 'NEGNUM',text)
    return re.sub(r'[^a-z0-9äöüß]','',text.lower().replace('\u00ad',''))
def main():
    parser=argparse.ArgumentParser();parser.add_argument('--volume',choices=['I','II'],required=True);parser.add_argument('--pdf',type=Path,required=True);parser.add_argument('--render-dir',type=Path,required=True);parser.add_argument('--toc-proof',type=Path);args=parser.parse_args()
    pdf=pymupdf.open(args.pdf);parts=[normalize(page.get_text()) for page in pdf];offsets=[0]
    for part in parts:offsets.append(offsets[-1]+len(part))
    full=''.join(parts);changes=json.loads((ROOT/'content/grundlagen/book-corrections-20261009.json').read_text())['entries'];changes=[e for e in changes if e['volume']==args.volume]
    docx=ROOT/f'source-assets/generated/books-20261009/PRINT_6x9_Band-{args.volume}_Konsistenz_2026-10-10.docx'
    doc=Document(docx);docx_sha=hashlib.sha256(docx.read_bytes()).hexdigest()
    source=next(x for x in json.loads((ROOT/'reports/consistency-20261009/book-source-build.json').read_text()) if x['volume']==args.volume)
    if source['outputSha256']!=docx_sha:raise RuntimeError('Current manuscript differs from its source-build proof')
    toc=json.loads(args.toc_proof.read_text()) if args.toc_proof else None
    signed_parts=[signed_normalize(page.get_text()) for page in pdf];signed_offsets=[0]
    for part in signed_parts:signed_offsets.append(signed_offsets[-1]+len(part))
    signed_full=''.join(signed_parts);table_cells={};table_proofs=[]
    for table_index in sorted({e['tableIndex'] for e in changes if 'tableIndex' in e}):
        table=doc.tables[table_index];header=''.join(signed_normalize(c.text) for c in table.rows[0].cells)
        flattened=[(ri,ci,signed_normalize(c.text)) for ri,row in enumerate(table.rows) for ci,c in enumerate(row.cells)]
        span=sum(len(value) for _,_,value in flattened)*2+300;starts=[];pos=signed_full.find(header)
        while pos>=0:starts.append(pos);pos=signed_full.find(header,pos+1)
        matches=[]
        for start in starts:
            cursor=start;cells={}
            for ri,ci,needle in flattened:
                pos=signed_full.find(needle,cursor,start+span) if needle else cursor
                if pos<0:break
                pages=list(range(bisect.bisect_right(signed_offsets,pos),bisect.bisect_left(signed_offsets,pos+len(needle))+1)) if needle else []
                cells[(ri,ci)]=pages;cursor=pos+len(needle)
            else:matches.append(cells)
        if len(matches)!=1:raise RuntimeError(f'Table {table_index} lacks a unique complete ordered-cell match in the rendered PDF')
        table_cells[table_index]=matches[0]
        table_proofs.append({'tableIndex':table_index,'orderedCellsVerified':len(matches[0]),'signedNumbersPreserved':True,'pages':sorted({p for pages in matches[0].values() for p in pages})})
    records=[];changed=set();unmatched=[]
    for entry in changes:
        actual=doc.paragraphs[entry['paragraphIndex']].text if 'paragraphIndex' in entry else doc.tables[entry['tableIndex']].rows[entry['rowIndex']].cells[entry['cellIndex']].text
        is_toc='paragraphIndex' in entry and doc.paragraphs[entry['paragraphIndex']].style.name.lower().startswith('toc')
        if is_toc:
            if actual.rsplit('\t',1)[0]!=entry['afterText'].rsplit('\t',1)[0] or not toc or toc['status']!='PASS' or toc['changedEntries']!=0 or toc['docxAfterSha256']!=docx_sha or toc['sourcePdfSha256']!=hashlib.sha256(args.pdf.read_bytes()).hexdigest():raise RuntimeError('TOC correction lacks a proof bound to this final render')
        elif signed_normalize(actual)!=signed_normalize(entry['afterText']):raise RuntimeError('Current corrected source text differs from the reviewed correction')
        needle=normalize(actual);pos=full.find(needle) if needle else -1;pages=[];occurrences=0
        if 'tableIndex' in entry:
            pages=table_cells[entry['tableIndex']][(entry['rowIndex'],entry['cellIndex'])];occurrences=1 if pages else 0;changed.update(pages)
        elif pos>=0:
            while pos>=0:
                pages.extend(range(bisect.bisect_right(offsets,pos),bisect.bisect_left(offsets,pos+len(needle))+1));occurrences+=1
                pos=full.find(needle,pos+len(needle))
            pages=sorted(set(pages));changed.update(pages)
        elif needle:unmatched.append(entry['paragraphIndex'] if 'paragraphIndex' in entry else f"table-{entry['tableIndex']}-{entry['rowIndex']}-{entry['cellIndex']}")
        records.append({'paragraphIndex':entry.get('paragraphIndex'),'tableIndex':entry.get('tableIndex'),'rowIndex':entry.get('rowIndex'),'cellIndex':entry.get('cellIndex'),'pages':pages,'occurrences':occurrences,'exactNormalizedTextLocated':bool(pages),'sourceTextVerified':True,'signedOrderedTableCellVerified':'tableIndex' in entry,'verifiedCurrentToc':is_toc,'deletedDuplicateSourceParagraph':not bool(needle)})
    captions=[]
    for p in doc.paragraphs:
        if re.match(r'^(Tabelle|Abbildung)\s+\d',p.text):
            caption=normalize(p.text);pages=[i+1 for i,part in enumerate(parts) if caption in part]
            captions.append({'caption':p.text,'pages':pages});changed.update(pages)
            # Inspect continued tables and both neighbouring page breaks.
            for page in pages:
                changed.update(n for n in [page-1,page+1] if 1<=n<=len(pdf))
    layout_proofs=[]
    relocated_breaks=[{'paragraphIndex':entry['pageBreakRelocatedToParagraphEnd'],
                       'imageDrawingPreserved':entry.get('imageDrawingPreserved',False)}
                     for entry in source.get('pageBreakLayout',[]) if 'pageBreakRelocatedToParagraphEnd' in entry]
    for entry in source.get('sdgBodyLayout',[])+source.get('shortContinuationLayout',[])+source.get('referenceBreakLayout',[])+relocated_breaks:
        index=entry['paragraphIndex'];needle=normalize(doc.paragraphs[index].text)
        if not needle and entry.get('imageDrawingPreserved'):
            paragraph=doc.paragraphs[index]
            if len(paragraph._p.xpath('.//w:drawing'))!=1 or len(paragraph._p.xpath('.//w:br[@w:type="page"]'))!=1:
                raise RuntimeError('Relocated image-page break lacks its preserved drawing and break')
            anchor=next((i for i in range(index-1,-1,-1) if normalize(doc.paragraphs[i].text)),None)
            if anchor is None:raise RuntimeError('Relocated image-page break lacks a textual anchor')
            anchor_text=normalize(doc.paragraphs[anchor].text)
            anchor_pages=[i+1 for i,part in enumerate(parts) if anchor_text in part and pdf[i].get_image_info()]
            if len(anchor_pages)!=1:raise RuntimeError('Relocated image-page break lacks a unique rendered figure page')
            pages=sorted({p for n in anchor_pages for p in (n-1,n,n+1) if 1<=p<=len(pdf)})
            changed.update(pages)
            layout_proofs.append({'paragraphIndex':index,'pages':pages,'imageDrawingPreserved':True,
                                  'renderedFigurePage':anchor_pages[0],'anchorParagraphIndex':anchor})
            continue
        position=full.find(needle)
        if not needle or position<0:raise RuntimeError('Layout-corrected paragraph is absent from the final PDF')
        pages=list(range(bisect.bisect_right(offsets,position),bisect.bisect_left(offsets,position+len(needle))+1))
        changed.update(pages);layout_proofs.append({'paragraphIndex':index,'pages':pages,'textLocated':True})
    outside=[];missing=[];images=0
    for i,page in enumerate(pdf):
        if not (args.render_dir/f'page-{i+1}.png').exists():missing.append(i+1)
        images+=len(page.get_image_info())
        # Also inspect every actual image page: a figure-only revision need not
        # change a paragraph or have a separately rendered text caption.
        if page.get_image_info():changed.add(i+1)
        for block in page.get_text('dict')['blocks']:
            rect=pymupdf.Rect(block['bbox'])
            if rect.width and rect.height and not (page.rect+(-1,-1,1,1)).contains(rect):outside.append({'page':i+1,'bbox':list(rect),'type':block['type']})
    report={'volume':args.volume,'pdf':args.pdf.name,'pdfSha256':hashlib.sha256(args.pdf.read_bytes()).hexdigest(),'docxSha256':docx_sha,'pages':len(pdf),'renderedPages':len(pdf)-len(missing),'missingPageRenders':missing,'sourceCorrections':len(changes),'locatedCorrections':sum(r['exactNormalizedTextLocated'] for r in records),'unmatchedCorrectionLocations':unmatched,'correctionLocations':records,'orderedTableProofs':table_proofs,'captions':captions,'layoutParagraphProofs':layout_proofs,'imagePlacements':images,'contentOutsidePage':outside,'pagesForDetailedVisualReview':sorted(changed),'visualReview':'PENDING'}
    target=ROOT/f'reports/consistency-20261009/book-render-audit-{args.volume}.json';target.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:report[k] for k in ['volume','pages','renderedPages','sourceCorrections','locatedCorrections','unmatchedCorrectionLocations','imagePlacements','contentOutsidePage']},ensure_ascii=False))
    if missing or outside or unmatched:raise SystemExit(1)
if __name__=='__main__':main()
