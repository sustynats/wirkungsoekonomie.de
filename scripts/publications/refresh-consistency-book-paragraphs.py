"""Apply a reviewed late correction without resetting the checked TOC or media.

The immutable manuscript and the generated file are checksum-bound. Only the
selected top-level paragraphs may change; all other XML and ZIP members retain
their content. This complements, rather than replaces, the full source builder.
"""
from pathlib import Path
from zipfile import ZipFile
from copy import deepcopy
import argparse, hashlib, json, os, re
from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
def sha(data):
    return hashlib.sha256(data).hexdigest()

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--volume', choices=['I', 'II'], required=True)
    parser.add_argument('--paragraph-index', type=int, action='append', required=True)
    args = parser.parse_args()
    report_path = ROOT / 'reports/consistency-20261009/book-source-build.json'
    reports = json.loads(report_path.read_text())
    record = next(r for r in reports if r['volume'] == args.volume)
    original = ROOT / record['source']
    output = ROOT / record['output']
    if sha(original.read_bytes()) != record['sourceSha256']:
        raise RuntimeError('Original manuscript changed')
    if sha(output.read_bytes()) != record['outputSha256']:
        raise RuntimeError('Unknown generated manuscript')
    entries = [e for e in json.loads((ROOT / 'content/grundlagen/book-corrections-20261009.json').read_text())['entries'] if e['volume'] == args.volume]
    wanted = set(args.paragraph_index)
    selected = [e for e in entries if e.get('paragraphIndex') in wanted]
    if len(selected) != len(wanted):
        raise RuntimeError('Missing or duplicate reviewed correction')
    source = Document(original)
    doc = Document(output)
    before_paragraphs = [etree.tostring(p._p) for p in doc.paragraphs]
    before_tables = [etree.tostring(t._tbl) for t in doc.tables]
    changed = []
    for e in selected:
        index = e['paragraphIndex']
        if source.paragraphs[index].text != e['beforeText'] or e['sourceSha256'] != record['sourceSha256']:
            raise RuntimeError('Correction does not match original source')
        paragraph = doc.paragraphs[index]
        after = re.sub('[\u2010\u2011\u2012\u2013\u2014\u2212]', '-', e['afterText']).replace('\u00ad', '')
        if paragraph.text == after:
            continue
        before = re.sub('[\u2010\u2011\u2012\u2013\u2014\u2212]', '-', e['beforeText']).replace('\u00ad', '')
        if paragraph.text != before:
            raise RuntimeError('Generated paragraph conflict')
        if paragraph._p.xpath('.//w:drawing | .//w:footnoteReference | .//w:endnoteReference'):
            raise RuntimeError('Embedded object requires separate review')
        if re.search(r'[+<>Σ∑×→≥≤]', after):
            raise RuntimeError('Mathematical typography requires the full builder')
        formatting = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
        paragraph.clear()
        run = paragraph.add_run(after)
        if formatting is not None:
            run._r.insert(0, formatting)
        changed.append(index)
    if before_tables != [etree.tostring(t._tbl) for t in doc.tables]:
        raise RuntimeError('A table changed')
    for index, paragraph in enumerate(doc.paragraphs):
        if index not in changed and etree.tostring(paragraph._p) != before_paragraphs[index]:
            raise RuntimeError('A non-target paragraph changed')
    with ZipFile(output) as archive:
        members = {n: archive.read(n) for n in archive.namelist()}
        info = archive.infolist()
    replacement = doc.part.blob
    before_sha = sha(output.read_bytes())
    if changed:
        temporary = output.with_suffix('.paragraphs.tmp')
        with ZipFile(temporary, 'w') as archive:
            for entry in info:
                archive.writestr(entry, replacement if entry.filename == 'word/document.xml' else members[entry.filename])
        with ZipFile(temporary) as archive:
            for name, data in members.items():
                if name != 'word/document.xml' and archive.read(name) != data:
                    raise RuntimeError('A non-target ZIP member changed')
        os.replace(temporary, output)
    record['outputSha256'] = sha(output.read_bytes())
    record['corrections'] = len(entries)
    record['finalParagraphRefresh'] = dict(beforeDocxSha256=before_sha, afterDocxSha256=record['outputSha256'],
        changedParagraphs=changed, preservedNonTargetMembers=len(members)-1,
        tablesTocMediaRelationshipsAndMetadataUnchanged=True)
    report_path.write_text(json.dumps(reports, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps(record['finalParagraphRefresh']))

if __name__ == '__main__':
    main()
