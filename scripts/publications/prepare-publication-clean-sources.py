#!/usr/bin/env python3
"""Create corrected publishing sources while preserving archived originals."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
ORIGINAL = ROOT / "outputs/nonpublic-source-archive/docs/oeffentlicher-wirkungsraum/source/woek_dossier_oeffentlicher_wirkungsraum_wellen_tiefe_v0_1.docx"
ORIGINAL_SHA256 = "9c801a8378a9e81b4bf24d0a568b327b89d7c930dcbb038f7a7baaaa257ea351"
OUTPUT = ROOT / "source-assets/generated/publication-cleanup-20261013/woek_dossier_oeffentlicher_wirkungsraum_wellen_tiefe_v0_1.docx"
REPORT = ROOT / "reports/consistency-20261009/publication-clean-source-report.json"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> None:
    if sha256(ORIGINAL) != ORIGINAL_SHA256:
        raise RuntimeError("The archived dossier source differs from the reviewed original")
    document = Document(ORIGINAL)
    replacements = {
        "28. Redaktionshinweis für die nächste Fassung": "28. Schlussfolgerung: Stein, Wellen und Tiefe",
        "Für die nächste Fassung sollte das Dossier redaktionell neu sortiert werden: Quellen und Bezugslinien gehören an das Ende; die Erweiterungskapitel sollten in die Hauptstruktur integriert werden; die Fallbibliothek kann als eigener Anhang stehen. Außerdem sollten ausgewählte Kapitel mit externen wissenschaftlichen Quellen, Medienforschung und Demokratietheorie stärker belegt werden.": (
            "Öffentliche Debatten lassen sich nicht allein über einzelne Aussagen verstehen. "
            "Die fünf Wellen zeigen, wie Aufmerksamkeit, Emotion, Deutung, soziale Verstärkung und politische Verschiebung zusammenwirken. "
            "Die Tiefenanalyse ergänzt diese Dynamik um Ursachen, Institutionen, materielle Bedingungen und demokratische Schutzräume."
        ),
        "Trotzdem ist die Grundarchitektur bereits tragfähig: Der öffentliche Wirkungsraum ergänzt die Wirkungsökonomie dort, wo Preise, Steuern und Produkte nicht ausreichen - im Raum der Aufmerksamkeit, Deutung, Resonanz und demokratischen Korrektur. Damit wird sichtbar, warum es nicht nur um Fakten geht und warum der Satz \"Handlungen haben Wirkung\" die Arbeit nicht beendet, sondern erst eröffnet.": (
            "Der öffentliche Wirkungsraum ergänzt die Wirkungsökonomie dort, wo Preise, Steuern und Produkte nicht ausreichen: "
            "im Raum der Aufmerksamkeit, Deutung, Resonanz und demokratischen Korrektur. Damit wird sichtbar, warum es nicht nur um Fakten geht "
            "und warum der Satz \"Handlungen haben Wirkung\" die Analyse nicht beendet, sondern erst eröffnet."
        ),
        "Die stärkste künftige Kurzformel lautet: Fakten prüfen den Stein. Wirkungsanalyse liest die Wellen. Wirkungsökonomie verändert die Tiefe.": (
            "Die Kurzformel lautet: Fakten prüfen den Stein. Wirkungsanalyse liest die Wellen. Wirkungsökonomie verändert die Tiefe."
        ),
        "27. Umsetzungspfad für ein 100-Seiten-Dossier, Website und Akademie": "27. Erprobung und institutionelle Verankerung",
    }
    deletions = {
        "Nächste Textprodukte",
        'Kurzmanifest "Öffentlicher Wirkungsraum"',
        "Glossar und Hoverdefinitionen für Stein, Welle, Tiefe, Resonanzraum, Wirkungsgewicht",
        "Debattenkarten als Website-Tool",
        "Akademiemodul mit Übungen",
        "Whitepaper zur demokratischen Sicherung",
        "Methodenpapier Wirkungsradar",
        "Serie mit 10 Fallstudien",
    }
    found = {key: 0 for key in replacements}
    for paragraph in document.paragraphs:
        if paragraph.text in replacements:
            found[paragraph.text] += 1
            paragraph.text = replacements[paragraph.text]
    if any(count != 1 for count in found.values()):
        raise RuntimeError(f"Expected each dossier correction exactly once: {found}")
    deleted = {value: 0 for value in deletions}
    for paragraph in list(document.paragraphs):
        if paragraph.text in deletions:
            deleted[paragraph.text] += 1
            paragraph._element.getparent().remove(paragraph._element)
    if any(count != 1 for count in deleted.values()):
        raise RuntimeError(f"Expected each production-only dossier block exactly once: {deleted}")
    document.core_properties.author = "Natalie Weber"
    document.core_properties.last_modified_by = "Institut für Wirkungsökonomie"
    document.core_properties.comments = ""
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    REPORT.parent.mkdir(parents=True, exist_ok=True)
    report = {
        "date": "2026-10-13",
        "status": "PASS",
        "original": str(ORIGINAL.relative_to(ROOT)),
        "originalSha256": ORIGINAL_SHA256,
        "publishingSource": str(OUTPUT.relative_to(ROOT)),
        "publishingSourceSha256": sha256(OUTPUT),
        "replacements": len(replacements),
        "productionBlocksRemoved": len(deletions),
        "preservation": "The archived original remains byte-identical; only the generated publishing source is revised.",
    }
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
