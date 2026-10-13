"""Protect embedded source figures against destructive paragraph replacement."""
from pathlib import Path
import importlib.util
import unittest
from io import BytesIO
from PIL import Image
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
spec = importlib.util.spec_from_file_location('book_source_edit', ROOT / 'scripts/publications/book_source_edit.py')
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


class BookSourceEdit(unittest.TestCase):
    def figure_paragraph(self):
        doc = Document(); paragraph = doc.add_paragraph()
        stream = BytesIO(); Image.new('RGB', (20, 10), 'green').save(stream, format='PNG'); stream.seek(0)
        paragraph.add_run().add_picture(stream)
        paragraph.add_run('Ungeprüfte alte Aussage.').bold = True
        return doc, paragraph

    def test_bound_figure_and_relationship_survive_corrected_text(self):
        doc, paragraph = self.figure_paragraph()
        drawing = paragraph._p.xpath('.//w:drawing')[0].xml
        rid = paragraph._p.xpath('.//a:blip')[0].get(qn('r:embed'))
        blob = doc.part.rels[rid].target_part.blob
        module.replace_paragraph(paragraph, 'Geprüfte Aussage - mit Quelle.', prefix_image='image1.png')
        self.assertEqual(paragraph.text, 'Geprüfte Aussage - mit Quelle.')
        self.assertEqual(paragraph._p.xpath('.//w:drawing')[0].xml, drawing)
        self.assertEqual(doc.part.rels[rid].target_part.blob, blob)
        self.assertTrue(paragraph.runs[-1].bold)
        saved = BytesIO(); doc.save(saved); saved.seek(0)
        self.assertEqual(Document(saved).paragraphs[0].text, paragraph.text)

    def test_unreviewed_or_wrong_figure_fails_before_mutation(self):
        _, paragraph = self.figure_paragraph(); before = paragraph._p.xml
        for expected in [None, 'image66.jpeg']:
            with self.assertRaises(RuntimeError):
                module.replace_paragraph(paragraph, 'Neue Aussage', prefix_image=expected)
            self.assertEqual(paragraph._p.xml, before)

    def test_unknown_object_and_note_reference_are_retained_on_failure(self):
        _, paragraph = self.figure_paragraph()
        paragraph.runs[0]._r.append(OxmlElement('w:tab')); before = paragraph._p.xml
        with self.assertRaises(RuntimeError): module.replace_paragraph(paragraph, 'Neue Aussage', prefix_image='image1.png')
        self.assertEqual(paragraph._p.xml, before)
        doc = Document(); paragraph = doc.add_paragraph('Quellenhinweis')
        note = OxmlElement('w:footnoteReference'); note.set(qn('w:id'), '17'); paragraph.runs[0]._r.append(note)
        before = paragraph._p.xml
        with self.assertRaises(RuntimeError): module.replace_paragraph(paragraph, 'Neue Aussage')
        self.assertEqual(paragraph._p.xml, before)


if __name__ == '__main__': unittest.main()
