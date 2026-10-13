"""Source-preservation regressions for the verified PDF ligature failure."""
from pathlib import Path
import importlib.util
import unittest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
spec = importlib.util.spec_from_file_location('book_typography', ROOT / 'scripts/publications/book_typography.py')
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


class BookTypography(unittest.TestCase):
    def test_words_tabs_breaks_bold_and_footnote_are_preserved(self):
        doc = Document()
        paragraph = doc.add_paragraph()
        run = paragraph.add_run('Regeln überflüssig\tKapitalflüsse\nFlächen bleiben.');run.bold = True
        note = OxmlElement('w:footnoteReference');note.set(qn('w:id'), '17');run._r.append(note)
        before = paragraph.text
        proof = module.format_umlaut_ligatures(doc)
        self.assertEqual(paragraph.text, before)
        self.assertEqual(proof['wordRuns'], 3)
        self.assertEqual(len(doc.element.xpath('.//w:footnoteReference')), 1)
        self.assertEqual(doc.element.xpath('.//w:footnoteReference')[0].get(qn('w:id')), '17')
        self.assertTrue(all(r.bold for r in paragraph.runs))
        self.assertEqual([r.text for r in paragraph.runs if r.font.name == 'DejaVu Serif'],
                         ['überflüssig', 'Kapitalflüsse', 'Flächen'])

    def test_unaffected_text_and_styles_are_byte_unchanged(self):
        doc = Document();doc.add_paragraph('Wirkung ist eine tatsächliche Zustandsveränderung.')
        before = doc.element.xml
        self.assertEqual(module.format_umlaut_ligatures(doc)['wordRuns'], 0)
        self.assertEqual(doc.element.xml, before)

    def test_repeat_application_preserves_all_source_text(self):
        doc = Document();paragraph = doc.add_paragraph('Flüsse und überflüssige Vorschriften.')
        before = paragraph.text
        module.format_umlaut_ligatures(doc);module.format_umlaut_ligatures(doc)
        self.assertEqual(paragraph.text, before)
        self.assertEqual(len([r for r in paragraph.runs if r.font.name == 'DejaVu Serif']), 2)


if __name__ == '__main__':
    unittest.main()
